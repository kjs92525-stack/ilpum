import openpyxl, re, json
wb=openpyxl.load_workbook('oct.xlsx')
ws=wb['10월 스케줄']
TITLE=re.compile(r'^(\d+)/(\d+) \((.)\)')
cells={}
for row in ws.iter_rows():
    for c in row:
        if c.value not in (None,''): cells[(c.row,c.column)]=c.value
titles=[(r,c,m) for (r,c),v in cells.items() if isinstance(v,str) and (m:=TITLE.match(v))]
titles.sort()
expected={}   # date -> {'pos':{posname:[tokens]}, 'off':[tokens]}
for r,c,m in titles:
    mo,d=int(m.group(1)),int(m.group(2)); date=f'2026-{mo:02d}-{d:02d}'
    pos={}; off=[]; rr=r+1
    maxr=max(k[0] for k in cells)
    while rr<=maxr:
        lab=cells.get((rr,c))
        if isinstance(lab,str) and TITLE.match(lab): break
        if lab is None: rr+=1; continue
        if isinstance(lab,str) and lab.startswith('휴무:'):
            off=[x.strip() for x in lab[3:].split(',') if x.strip()]; break
        names=cells.get((rr,c+1),'')
        pos[lab]=[x.strip() for x in str(names).split(', ') if x.strip()] if names else []
        rr+=1
    expected[date]={'pos':pos,'off':off,'title':cells[(r,c)]}
json.dump(expected,open('expected.json','w'),ensure_ascii=False,indent=1)
print(len(expected),'days', min(expected),max(expected))
for d in ['2026-10-01','2026-10-12','2026-10-31']:
    print(d, expected[d]['title']); 
    for p,n in expected[d]['pos'].items(): print('  ',p,n)
    print('   휴무',expected[d]['off'])
# 근무자
sw=wb['근무자']; staff=[]
for r in sw.iter_rows(min_row=2,values_only=True):
    if r[0]: staff.append(dict(raw=r[0],pos=r[1],off=r[2],days=r[3]))
json.dump(staff,open('staff.json','w'),ensure_ascii=False,indent=1)
ew=wb['예외']; ex=[dict(date=r[0],type=r[1],name=r[2],pos=r[3] or '',memo=r[4] or '') for r in ew.iter_rows(min_row=2,values_only=True) if r[0]]
json.dump(ex,open('ex.json','w'),ensure_ascii=False,indent=1)
mw=wb['메모']; memo=[dict(date=r[0],name=r[1],kind=r[2],memo=r[3]) for r in mw.iter_rows(min_row=2,values_only=True) if r[0]]
json.dump(memo,open('memo.json','w'),ensure_ascii=False,indent=1)
print(len(staff),'staff',len(ex),'ex',len(memo),'memo')
